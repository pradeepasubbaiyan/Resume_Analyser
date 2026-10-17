from __future__ import annotations

from datetime import datetime, timezone
from functools import lru_cache
import ipaddress
import http.client
import re
import socket
import ssl
from html.parser import HTMLParser
from io import BytesIO
from pathlib import PurePath
import zipfile
from urllib.parse import quote, urlsplit
from xml.etree import ElementTree

import fitz
import httpx
from docx import Document
from fastapi import HTTPException, UploadFile

from app.config import get_settings
from app.workflow import normalize_role, skill_match, skills_for_role, weighted_role_fit

MAX_RESUME_BYTES = 5 * 1024 * 1024
MAX_PORTFOLIO_BYTES = 1024 * 1024
MAX_GITHUB_REPOSITORIES = 20
MAX_GITHUB_READMES = 5
MAX_GITHUB_SOURCE_BYTES = 150_000
MAX_RESUME_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_OCR_IMAGE_PIXELS = 25_000_000
MAX_OCR_PDF_PAGES = 20
IMAGE_SUFFIXES = {'.bmp', '.gif', '.jpeg', '.jpg', '.png', '.tif', '.tiff', '.webp'}


class _VisibleTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {'script', 'style', 'noscript', 'svg'}:
            self.hidden_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {'script', 'style', 'noscript', 'svg'} and self.hidden_depth:
            self.hidden_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self.hidden_depth:
            text = data.strip()
            if text:
                self.parts.append(text)


def _extract_resume(filename: str, contents: bytes) -> str:
    if len(contents) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail='Resume file must be no larger than 5 MB')

    suffix = PurePath(filename).suffix.lower()
    if contents.startswith(b'%PDF-'):
        suffix = '.pdf'
    elif contents.startswith(b'{\\rtf'):
        suffix = '.rtf'
    elif contents.startswith(b'PK\x03\x04'):
        try:
            with zipfile.ZipFile(BytesIO(contents)) as archive:
                names = set(archive.namelist())
            if 'word/document.xml' in names:
                suffix = '.docx'
            elif 'content.xml' in names:
                suffix = '.odt'
            elif any(name.startswith('ppt/slides/slide') and name.endswith('.xml') for name in names):
                suffix = '.pptx'
            elif 'xl/workbook.xml' in names:
                suffix = '.xlsx'
        except zipfile.BadZipFile:
            pass
    try:
        if suffix == '.pdf':
            with fitz.open(stream=contents, filetype='pdf') as document:
                if len(document) > 100:
                    raise HTTPException(status_code=413, detail='Resume PDF cannot exceed 100 pages')
                text = '\n'.join(page.get_text() for page in document)
                if not text.strip():
                    if len(document) > MAX_OCR_PDF_PAGES:
                        raise HTTPException(
                            status_code=413,
                            detail=f'Scanned PDFs can contain no more than {MAX_OCR_PDF_PAGES} pages for OCR.',
                        )
                    text = '\n'.join(
                        _extract_image_text(page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes('png'))
                        for page in document
                    )
        elif suffix == '.docx':
            with zipfile.ZipFile(BytesIO(contents)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_RESUME_UNCOMPRESSED_BYTES:
                    raise HTTPException(status_code=413, detail='Expanded Word document exceeds the 20 MB extraction limit')
            document = Document(BytesIO(contents))
            text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
            text += '\n' + '\n'.join(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
        elif suffix == '.odt':
            with zipfile.ZipFile(BytesIO(contents)) as archive:
                content = archive.getinfo('content.xml')
                if content.file_size > MAX_RESUME_UNCOMPRESSED_BYTES:
                    raise HTTPException(status_code=413, detail='Expanded OpenDocument exceeds the 20 MB extraction limit')
                root = ElementTree.fromstring(archive.read(content))
            text = '\n'.join(value.strip() for value in root.itertext() if value.strip())
        elif suffix == '.rtf':
            decoded = contents.decode('cp1252', errors='replace')
            decoded = re.sub(r"\{\\\*[^{}]*\}", ' ', decoded)
            decoded = re.sub(r"\\'[0-9a-fA-F]{2}", lambda match: bytes.fromhex(match.group(0)[2:]).decode('cp1252'), decoded)
            decoded = re.sub(r'\\u(-?\d+)\??', lambda match: chr(int(match.group(1)) % 65536), decoded)
            decoded = re.sub(r'\\(?:par|line|tab)\b ?', '\n', decoded)
            decoded = re.sub(r'\\[a-zA-Z]+-?\d* ?', '', decoded)
            text = decoded.replace('{', ' ').replace('}', ' ')
        elif suffix == '.pptx':
            text = _extract_openxml_text(contents, 'ppt/slides/')
        elif suffix == '.xlsx':
            text = _extract_openxml_text(contents, 'xl/')
        elif suffix in IMAGE_SUFFIXES or not suffix:
            text = _extract_image_text(contents)
        elif suffix in {'.txt', '.md', '.rst', '.tex', '.csv', '.json', '.yaml', '.yml', '.html', '.htm', '.xml'}:
            text = _decode_text(contents)
            if suffix in {'.html', '.htm', '.xml'}:
                parser = _VisibleTextParser()
                parser.feed(text)
                text = '\n'.join(parser.parts)
        else:
            try:
                text = _decode_text(contents)
            except HTTPException as error:
                raise HTTPException(
                    status_code=415,
                    detail=(
                        f'The {suffix or "uploaded"} file format is not supported for text extraction. '
                        'Use PDF, Word (.docx), OpenDocument (.odt), RTF, PowerPoint (.pptx), '
                        'Excel (.xlsx), an image, or a text-readable file.'
                    ),
                ) from error
            if not _looks_like_text(text):
                raise HTTPException(
                    status_code=415,
                    detail=(
                        f'The {suffix or "uploaded"} file format is not supported for text extraction. '
                        'Use PDF, Word (.docx), OpenDocument (.odt), RTF, PowerPoint (.pptx), '
                        'Excel (.xlsx), an image, or a text-readable file.'
                    ),
                )
    except HTTPException:
        raise
    except (KeyError, zipfile.BadZipFile, ElementTree.ParseError) as error:
        raise HTTPException(status_code=400, detail='The uploaded document is malformed or unreadable') from error
    except Exception as error:
        raise HTTPException(status_code=400, detail='Could not read the uploaded resume file') from error

    if not text.strip():
        raise HTTPException(
            status_code=422,
            detail='No readable text was found in the resume. For image-based documents, upload a clear scan with legible text.',
        )
    return text[:200_000]


def _extract_openxml_text(contents: bytes, prefix: str) -> str:
    parts: list[str] = []
    with zipfile.ZipFile(BytesIO(contents)) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_RESUME_UNCOMPRESSED_BYTES:
            raise HTTPException(status_code=413, detail='Expanded Office document exceeds the 20 MB extraction limit')
        for item in archive.infolist():
            if not item.filename.startswith(prefix) or not item.filename.endswith('.xml'):
                continue
            root = ElementTree.fromstring(archive.read(item))
            parts.extend(value.strip() for value in root.itertext() if value.strip())
    return '\n'.join(parts)


@lru_cache(maxsize=1)
def _get_ocr_engine():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR()


def _extract_image_text(contents: bytes) -> str:
    try:
        import numpy
        from PIL import Image, UnidentifiedImageError

        with Image.open(BytesIO(contents)) as image:
            if image.width * image.height > MAX_OCR_IMAGE_PIXELS:
                raise HTTPException(status_code=413, detail='Resume images cannot exceed 25 megapixels')
            image.seek(0)
            rgb_image = image.convert('RGB')
            if max(rgb_image.size) > 3000:
                rgb_image.thumbnail((3000, 3000))
            result, _ = _get_ocr_engine()(numpy.asarray(rgb_image))
    except HTTPException:
        raise
    except UnidentifiedImageError as error:
        raise HTTPException(
            status_code=415,
            detail='The image format is not supported or the uploaded file is not a valid image.',
        ) from error
    except Exception as error:
        raise HTTPException(status_code=422, detail='Could not read text from the uploaded image') from error

    if not result:
        return ''
    return '\n'.join(str(line[1]) for line in result if len(line) > 1 and line[1])


def _decode_text(contents: bytes) -> str:
    for encoding in ('utf-8-sig', 'utf-16', 'cp1252'):
        try:
            decoded = contents.decode(encoding)
            if _looks_like_text(decoded):
                return decoded
        except UnicodeDecodeError:
            continue
    raise HTTPException(status_code=415, detail='The uploaded file does not contain readable text')


def _looks_like_text(text: str) -> bool:
    if not text:
        return False
    printable = sum(character.isprintable() or character in '\r\n\t' for character in text)
    return '\x00' not in text and printable / len(text) >= 0.85


def _github_username(profile_url: str) -> str:
    parsed = urlsplit(profile_url)
    if parsed.scheme != 'https' or parsed.hostname != 'github.com' or parsed.username or parsed.password:
        raise HTTPException(status_code=422, detail='Use a public https://github.com/username profile URL')
    segments = [segment for segment in parsed.path.split('/') if segment]
    if len(segments) != 1 or not re.fullmatch(r'[A-Za-z0-9-]{1,39}', segments[0]):
        raise HTTPException(status_code=422, detail='GitHub URL must point to a profile, not a repository')
    return segments[0]


def _validated_portfolio_url(portfolio_url: str) -> tuple[str, str, str]:
    parsed = urlsplit(portfolio_url)
    hostname = parsed.hostname
    if parsed.scheme != 'https' or not hostname or parsed.username or parsed.password:
        raise HTTPException(status_code=422, detail='Portfolio URL must be a public HTTPS URL')
    try:
        port = parsed.port
    except ValueError as error:
        raise HTTPException(status_code=422, detail='Portfolio URL contains an invalid port') from error
    if port not in (None, 443):
        raise HTTPException(status_code=422, detail='Portfolio URLs on non-standard ports cannot be analyzed')
    if hostname.lower() in {'localhost', 'localhost.localdomain'}:
        raise HTTPException(status_code=422, detail='Local and private portfolio URLs cannot be analyzed')
    try:
        addresses = socket.getaddrinfo(hostname, parsed.port or 443, type=socket.SOCK_STREAM)
    except OSError as error:
        raise HTTPException(status_code=422, detail='Portfolio host could not be resolved') from error
    if not addresses or any(
        not ipaddress.ip_address(address[4][0]).is_global for address in addresses
    ):
        raise HTTPException(status_code=422, detail='Portfolio URL must resolve only to public IP addresses')
    return parsed.geturl(), hostname, addresses[0][4][0]


def _fetch_github(username: str) -> tuple[str, dict[str, str]]:
    headers = {
        'Accept': 'application/vnd.github+json',
        'User-Agent': 'VerityHire-local-evidence-review',
    }
    github_token = get_settings().github_token
    if github_token:
        headers['Authorization'] = f'Bearer {github_token}'
    try:
        with httpx.Client(timeout=8, follow_redirects=False, trust_env=False) as client:
            profile_response = client.get(f'https://api.github.com/users/{username}', headers=headers)
            if profile_response.status_code == 404:
                return '', {'status': 'not found', 'detail': 'Public GitHub profile was not found'}
            if not profile_response.is_success:
                return '', {'status': 'unavailable', 'detail': _github_error_detail(profile_response)}
            profile = profile_response.json()
            repositories_response = client.get(
                f'https://api.github.com/users/{username}/repos',
                params={'type': 'owner', 'sort': 'updated', 'per_page': MAX_GITHUB_REPOSITORIES},
                headers=headers,
            )
            if not repositories_response.is_success:
                return '', {'status': 'unavailable', 'detail': _github_error_detail(repositories_response)}
            repositories = repositories_response.json()
            pieces = [profile.get('bio') or '', profile.get('name') or '']
            for repository in repositories:
                pieces.extend([
                    repository.get('name') or '',
                    repository.get('description') or '',
                    repository.get('language') or '',
                    ' '.join(repository.get('topics') or []),
                ])
            readme_count = 0
            for repository in repositories[:MAX_GITHUB_READMES]:
                repository_name = repository.get('name')
                if not repository_name:
                    continue
                try:
                    readme_response = client.get(
                        f'https://api.github.com/repos/{username}/{repository_name}/readme',
                        headers={**headers, 'Accept': 'application/vnd.github.raw'},
                    )
                    if readme_response.is_success:
                        pieces.append(f'Repository {repository_name} README:\n{readme_response.text[:50_000]}')
                        readme_count += 1
                    elif _is_github_rate_limited(readme_response):
                        return '\n'.join(pieces), {
                            'status': 'partial',
                            'detail': _github_error_detail(readme_response),
                        }
                except httpx.HTTPError:
                    continue
                branch = repository.get('default_branch')
                if not branch:
                    continue
                try:
                    tree_response = client.get(
                        f'https://api.github.com/repos/{username}/{repository_name}/git/trees/{quote(branch, safe="")}',
                        params={'recursive': '1'},
                        headers=headers,
                    )
                    if not tree_response.is_success:
                        if _is_github_rate_limited(tree_response):
                            return '\n'.join(pieces), {
                                'status': 'partial',
                                'detail': _github_error_detail(tree_response),
                            }
                        continue
                    tree = tree_response.json().get('tree', [])
                    source_paths = [
                        item.get('path', '')
                        for item in tree
                        if item.get('type') == 'blob'
                        and item.get('size', MAX_GITHUB_SOURCE_BYTES + 1) <= 30_000
                        and re.search(
                            r'(^|/)(test[^/]*|[^/]*test[^/]*|package\.json|pyproject\.toml|requirements[^/]*\.txt|pom\.xml|go\.mod|Cargo\.toml)$|'
                            r'\.(py|ts|tsx|js|jsx|java|go|rs|sql|yaml|yml)$',
                            item.get('path', ''),
                            re.I,
                        )
                    ][:4]
                    for path in source_paths:
                        if sum(len(piece) for piece in pieces) >= MAX_GITHUB_SOURCE_BYTES:
                            break
                        try:
                            source_response = client.get(
                                f'https://api.github.com/repos/{username}/{repository_name}/contents/{quote(path, safe="/")}',
                                params={'ref': branch},
                                headers={**headers, 'Accept': 'application/vnd.github.raw'},
                            )
                            if source_response.is_success:
                                pieces.append(
                                    f'Repository {repository_name} file {path}:\n'
                                    f'{source_response.text[:30_000]}'
                                )
                            elif _is_github_rate_limited(source_response):
                                return '\n'.join(pieces), {
                                    'status': 'partial',
                                    'detail': _github_error_detail(source_response),
                                }
                        except httpx.HTTPError:
                            continue
                except (httpx.HTTPError, ValueError):
                    continue
    except httpx.TimeoutException:
        return '', {'status': 'unavailable', 'detail': 'GitHub API request timed out; retry later'}
    except httpx.HTTPError:
        return '', {'status': 'unavailable', 'detail': 'GitHub could not be read; check connectivity'}
    except ValueError:
        return '', {'status': 'unavailable', 'detail': 'GitHub returned an invalid API response'}

    return '\n'.join(pieces), {
        'status': 'analyzed',
        'detail': (
            f"Public profile, {len(repositories)} recently updated repositories, "
            f"and {readme_count} repository README files reviewed"
        ),
    }


def _is_github_rate_limited(response: httpx.Response) -> bool:
    return response.status_code == 403 and response.headers.get('x-ratelimit-remaining') == '0'


def _github_error_detail(response: httpx.Response) -> str:
    if response.status_code == 401:
        return 'GitHub rejected the server API token; check the GITHUB_TOKEN configuration'
    if _is_github_rate_limited(response):
        reset = response.headers.get('x-ratelimit-reset')
        if reset:
            try:
                reset_time = datetime.fromtimestamp(int(reset), timezone.utc).strftime('%Y-%m-%d %H:%M UTC')
                return f'GitHub API rate limit is exhausted; retry after {reset_time} or configure a server GITHUB_TOKEN'
            except (ValueError, OverflowError, OSError):
                pass
        return 'GitHub API rate limit is exhausted; retry later or configure a server GITHUB_TOKEN'
    if response.status_code == 403:
        return 'GitHub denied access to the public profile or repositories'
    if response.status_code >= 500:
        return f'GitHub API is temporarily unavailable (HTTP {response.status_code}); retry later'
    return f'GitHub API request failed with HTTP {response.status_code}'


def _fetch_portfolio(portfolio_url: str) -> tuple[str, dict[str, str]]:
    safe_url, hostname, resolved_ip = _validated_portfolio_url(portfolio_url)
    parsed = urlsplit(safe_url)
    path = parsed.path or '/'
    if parsed.query:
        path += f'?{parsed.query}'
    connection = None
    try:
        connection = _PinnedHTTPSConnection(hostname, resolved_ip)
        connection.request(
            'GET',
            path,
            headers={
                'Host': hostname,
                'User-Agent': 'VerityHire-local-evidence-review',
                'Accept-Encoding': 'identity',
            },
        )
        response = connection.getresponse()
        if 300 <= response.status < 400:
            return '', {'status': 'not analyzed', 'detail': 'Portfolio redirects are not followed'}
        if not 200 <= response.status < 300:
            return '', {'status': 'unavailable', 'detail': f'Portfolio request returned HTTP {response.status}'}
        content_type = response.getheader('content-type', '').lower()
        if 'text/html' not in content_type and 'text/plain' not in content_type:
            return '', {'status': 'not analyzed', 'detail': 'Portfolio must serve an HTML or plain-text page'}
        body = bytearray()
        while chunk := response.read(64 * 1024):
            body.extend(chunk)
            if len(body) > MAX_PORTFOLIO_BYTES:
                return '', {'status': 'not analyzed', 'detail': 'Portfolio page exceeds the 1 MB analysis limit'}
    except (OSError, ssl.SSLError, http.client.HTTPException):
        return '', {'status': 'unavailable', 'detail': 'Portfolio page could not be fetched'}
    finally:
        if connection:
            connection.close()

    parser = _VisibleTextParser()
    parser.feed(bytes(body).decode('utf-8', errors='replace'))
    return '\n'.join(parser.parts)[:100_000], {'status': 'analyzed', 'detail': 'Public portfolio page text reviewed'}


class _PinnedHTTPSConnection(http.client.HTTPSConnection):
    def __init__(self, hostname: str, resolved_ip: str) -> None:
        super().__init__(hostname, port=443, timeout=8, context=ssl.create_default_context())
        self.resolved_ip = resolved_ip

    def connect(self) -> None:
        raw_socket = socket.create_connection((self.resolved_ip, self.port), self.timeout)
        self.sock = self._context.wrap_socket(raw_socket, server_hostname=self.host)


def _contains_skill(text: str, skill: str) -> bool:
    return skill_match(text, skill)


async def analyze_candidate_evidence(
    *,
    target_role: str,
    role_description: str,
    resume_file: UploadFile | None,
    github_url: str | None,
    portfolio_url: str | None,
    role_weights: dict[str, float] | None = None,
) -> dict:
    try:
        target_role = normalize_role(target_role)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    if resume_file is None or not resume_file.filename:
        raise HTTPException(status_code=422, detail='A readable resume file is required')
    if not github_url or not github_url.strip():
        raise HTTPException(status_code=422, detail='A public GitHub profile URL is required')
    if len(role_description) > 4000:
        raise HTTPException(status_code=422, detail='Job description must be 4,000 characters or fewer')

    evidence: dict[str, str] = {}
    sources: dict[str, dict[str, str]] = {
        'resume': {'status': 'not provided', 'detail': 'No resume uploaded'},
        'github': {'status': 'not provided', 'detail': 'No GitHub profile provided'},
        'portfolio': {'status': 'not provided', 'detail': 'No portfolio URL provided'},
    }

    if resume_file is not None and resume_file.filename:
        resume_text = _extract_resume(resume_file.filename, await resume_file.read(MAX_RESUME_BYTES + 1))
        evidence['resume'] = resume_text
        sources['resume'] = {'status': 'analyzed', 'detail': f'{resume_file.filename} text extracted'}

    if github_url:
        username = _github_username(github_url)
        github_text, sources['github'] = _fetch_github(username)
        if github_text:
            evidence['github'] = github_text

    if portfolio_url:
        portfolio_text, sources['portfolio'] = _fetch_portfolio(portfolio_url)
        if portfolio_text:
            evidence['portfolio'] = portfolio_text

    if not evidence:
        raise HTTPException(
            status_code=422,
            detail='No resources could be analyzed. Check the resume file, GitHub profile, and portfolio accessibility.',
        )

    skills = skills_for_role(target_role, role_description)
    fit = weighted_role_fit(skills, evidence, role_weights, role_description)
    skill_evidence = []
    for skill in skills:
        found_in = [source for source, text in evidence.items() if _contains_skill(text, skill)]
        snippets = []
        for source in found_in:
            matching_lines = [
                (line_number, line.strip())
                for line_number, line in enumerate(evidence[source].splitlines(), 1)
                if _contains_skill(line, skill) and line.strip()
            ][:2]
            for line_number, line in matching_lines:
                snippets.append({
                    'source': source,
                    'source_location': f'extracted text line {line_number}' if source == 'resume' else f'public {source} text line {line_number}',
                    'excerpt': line[:400],
                })
        skill_evidence.append({
            'name': skill,
            'status': 'evidence found' if found_in else 'not found in analyzed sources',
            'result_category': 'supported by submitted evidence' if found_in else 'insufficient evidence',
            'sources': found_in,
            'evidence_snippets': snippets,
        })

    return {
        'data': {
            'profile': {'target_role': target_role},
            'report': {
                'score': fit['score'],
                'score_label': 'Evidence-based role alignment estimate',
                'is_probability': False,
                'analysis_method': 'role-specific evidence matching across the submitted resume, public GitHub data, and optional portfolio',
                'evidence_coverage': {
                    'resume_characters_analyzed': len(evidence.get('resume', '')),
                    'github_role_competencies_found': fit['github_metrics']['role_competencies_found'],
                    'github_repository_records_reviewed': fit['github_metrics']['repository_records_reviewed'],
                    'github_readme_files_reviewed': fit['github_metrics']['readme_files_reviewed'],
                    'github_source_files_reviewed': fit['github_metrics']['source_files_reviewed'],
                    'role_description_provided': bool(role_description.strip()),
                    'assessed_weight_percent': fit['assessed_weight_percent'],
                },
                'summary': (
                    f'Estimated alignment is {fit["score"]}/100 based on the weighted categories that could be '
                    f'assessed ({fit["assessed_weight_percent"]}% of configured weight). This is not a probability '
                    'of job success and does not establish proficiency.'
                ),
                'matched_skills': fit['matched_skills'],
                'skills_needing_verification': fit['skills_needing_verification'],
                'skill_evidence': skill_evidence,
                'role_framework': {
                    'competencies': skills,
                    'source': 'candidate-provided job description' if role_description.strip() else 'provisional role-title inference',
                    'assumptions': role_description.strip() or f'Competencies were inferred from the entered title "{target_role}" and may require recruiter confirmation.',
                },
                'weighting': fit['weights'],
                'score_categories': fit['categories'],
                'github_metrics': fit['github_metrics'],
                'assessed_weight_percent': fit['assessed_weight_percent'],
                'sources': sources,
                'supporting_evidence': [
                    f'{skill["name"]}: found in {", ".join(skill["sources"])}'
                    for skill in skill_evidence
                    if skill['sources']
                ],
                'limitations': [
                    'Role competencies are generated from deterministic role mappings and any supplied job description; no generative AI provider is configured.',
                    'Keyword matches point to evidence locations but do not verify proficiency or work ownership.',
                    'Public GitHub and portfolio data may be incomplete or rate-limited.',
                    'This report is not a hiring decision, success probability, or automatic rejection signal.',
                ],
            },
        },
        'message': 'Evidence analysis completed',
    }
